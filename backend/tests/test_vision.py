import base64

from conftest import FakeChatClient, make_pdf

from src.core import vision
from src.core.vision import Page, VisionProvider, read_document, read_pages, resolve_vision_provider


# ---------- provider resolution ----------
def test_no_provider_when_no_keys():
    assert resolve_vision_provider({}) is None


def test_vision_api_key_wins_over_everything():
    p = resolve_vision_provider({
        "VISION_API_KEY": "v", "VISION_BASE_URL": "https://example.test/v1", "VISION_MODEL": "my-vl",
        "OPENROUTER_API_KEY": "r", "OPENAI_API_KEY": "o",
    })
    assert (p.name, p.api_key, p.base_url, p.model) == ("custom", "v", "https://example.test/v1", "my-vl")


def test_openrouter_before_openai_with_defaults():
    p = resolve_vision_provider({"OPENROUTER_API_KEY": "r", "OPENAI_API_KEY": "o"})
    assert (p.name, p.base_url, p.model) == ("openrouter", "https://openrouter.ai/api/v1", "qwen/qwen2.5-vl-72b-instruct")


def test_openai_fallback_defaults():
    p = resolve_vision_provider({"OPENAI_API_KEY": "o"})
    assert (p.name, p.base_url, p.model) == ("openai", None, "gpt-4o-mini")


def test_resolution_reads_process_env(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "r")
    assert resolve_vision_provider().name == "openrouter"


# ---------- text-layer fallback ----------
def test_pdf_text_layer_fallback_without_key(tmp_path, monkeypatch, caplog):
    pdf = make_pdf(tmp_path / "paper.pdf", ["1. Solve x + 2 = 5. [2 marks]", "2. Factorise x^2 - 1. [3 marks]"])
    monkeypatch.setattr(vision, "make_vision_client", lambda *a, **k: (_ for _ in ()).throw(AssertionError("no client")))
    with caplog.at_level("WARNING"):
        pages = read_document(pdf, tmp_path / "work")
    assert "1. Solve x + 2 = 5. [2 marks]" in pages[0]
    assert "2. Factorise x^2 - 1. [3 marks]" in pages[1]
    assert "No vision provider configured" in caplog.text


def test_docx_text_layer_fallback_without_key(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Question 1 (a) Differentiate sin x. [2]")
    path = tmp_path / "paper.docx"
    doc.save(str(path))
    assert read_document(path, tmp_path / "work") == ["Question 1 (a) Differentiate sin x. [2]"]


# ---------- vision path (mocked client) ----------
PROVIDER = VisionProvider("openrouter", "test-key", "https://openrouter.ai/api/v1", "qwen/qwen2.5-vl-72b-instruct")
LATEX = "1. Solve $x^2 - 5x + 6 = 0$. [3 marks]"


def test_vision_path_transcribes_pdf_pages(tmp_path):
    pdf = make_pdf(tmp_path / "paper.pdf", ["", "scanned"])  # page 1 has no text layer at all
    client = FakeChatClient(lambda kw: LATEX)

    pages = read_document(pdf, tmp_path / "work", dpi=120, provider=PROVIDER, client=client)

    assert pages == [LATEX, LATEX]
    assert len(client.calls) == 2
    call = client.calls[0]
    assert call["model"] == "qwen/qwen2.5-vl-72b-instruct"
    assert "LaTeX" in call["messages"][0]["content"] and "marks" in call["messages"][0]["content"]
    url = next(p["image_url"]["url"] for p in call["messages"][1]["content"] if p["type"] == "image_url")
    assert base64.b64decode(url.split(",", 1)[1])[:4] == b"\x89PNG"


def test_docx_vision_uses_embedded_pictures(tmp_path, monkeypatch):
    import io

    from docx import Document
    from PIL import Image

    # Full-page DOCX rendering needs Word/LibreOffice; force the embedded-picture path.
    import src.core.render as render
    monkeypatch.setattr(render, "render_paper_to_images", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no soffice")))

    buf = io.BytesIO()
    Image.new("RGB", (200, 120), "white").save(buf, format="PNG")
    buf.seek(0)
    doc = Document()
    doc.add_paragraph("Q1. Evaluate the integral shown.")
    doc.add_picture(buf)
    path = tmp_path / "paper.docx"
    doc.save(str(path))

    client = FakeChatClient(lambda kw: "Q1. Evaluate $\\int_0^1 x\\,dx$. [2]")
    pages = read_document(path, tmp_path / "work", provider=PROVIDER, client=client)
    assert pages == ["Q1. Evaluate $\\int_0^1 x\\,dx$. [2]"]
    parts = client.calls[0]["messages"][1]["content"]
    assert sum(p["type"] == "image_url" for p in parts) == 1
    assert "Q1. Evaluate the integral shown." in parts[0]["text"]


def test_failed_page_falls_back_to_text_layer(monkeypatch):
    monkeypatch.setenv("VISION_MAX_RETRIES", "1")
    calls = {"n": 0}

    def always_fail(kw):
        calls["n"] += 1
        raise ConnectionError("simulated outage")

    out = read_pages([Page(1, "1. Native text [4]", [b"png"])], PROVIDER, client=FakeChatClient(always_fail))
    assert out == ["1. Native text [4]"]
    assert calls["n"] == 2


def test_empty_vision_response_falls_back(monkeypatch):
    monkeypatch.setenv("VISION_MAX_RETRIES", "0")
    out = read_pages([Page(1, "layer", [b"png"])], PROVIDER, client=FakeChatClient(lambda kw: "   "))
    assert out == ["layer"]


def test_page_cap(tmp_path, monkeypatch):
    monkeypatch.setenv("VISION_MAX_PAGES", "1")
    pdf = make_pdf(tmp_path / "paper.pdf", ["p1", "p2"])
    client = FakeChatClient(lambda kw: LATEX)
    pages = read_document(pdf, tmp_path / "work", dpi=120, provider=PROVIDER, client=client)
    assert pages[0] == LATEX and "p2" in pages[1]
    assert len(client.calls) == 1
